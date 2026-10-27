# -*- coding: utf-8 -*-
r"""test_api_commission.py — 委托手续套件接口回归（批 64，2026-10-13）

覆盖：
  · meta（6 件模板就位 + 档位选项 + 默认值）
  · plan（dry-run：必填 client 拦截 / 审级与授权方式档位校验 / 6 件预览 /
    cause_d 去「纠纷」尾 / payee 缺省=client）
  · build（confirm 门 dry-run / apply 六件生成到案卷 01-委托手续 / 重名避让 -2）
  · 生成读回：1 号分段结构 / 2 号授权书 / 3·5·6 表格 / 无 {{ 残留 /
    旧样例数据零残留（消毒）/ 特殊字符不双重转义
  · 安全闸：未选案卷 / 越界夹名 / 不存在夹

⚠️ 全隔离：案件根打桩到临时目录；**不碰真案卷**（只读模板、写临时目录）。
"""
import os
import re
import shutil
import sys
import tempfile
import zipfile

TMP = tempfile.mkdtemp(prefix='wb_comm_test_')
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

fails, total = [], []


def ck(name, cond):
    total.append(name)
    print(('  PASS ' if cond else '  FAIL ') + name)
    if not cond:
        fails.append(name)


def main():
    import app as A
    import docx

    tmp_root = os.path.join(TMP, 'cases')
    case = '执-测试案'
    os.makedirs(os.path.join(tmp_root, case))
    orig_root = A.case_ops.case_root
    A.case_ops.case_root = lambda: tmp_root
    try:
        # ---------- meta ----------
        m = A.api_commission_meta()
        ck('meta 六件模板全就位', m['ok'] and len(m['templates']) == 6)
        ck('meta 档位（审级 3 / 授权 2 / 类型 5）',
           len(m['stages']) == 3 and len(m['auth_types']) == 2 and len(m['case_types']) == 5)
        ck('meta 默认律师与法院', m['defaults']['lawyer'] == '安清惠'
           and '集美' in m['defaults']['court'])

        # ---------- plan ----------
        fields = {'client': '测试委托人甲', 'opponent': '测试对方乙', 'cause': '买卖合同纠纷',
                  'stage': '一', 'fee': '6000', 'auth_type': '特别', 'case_type': '民事',
                  'address': '测试地址', 'phone': '13800000000'}
        p = A.api_commission_plan({'fields': fields, 'case': case})
        ck('plan 六件预览', (not p.get('blocked')) and len(p['docs']) == 6)
        ck('plan cause_d 去「纠纷」尾', p['context']['cause_d'] == '买卖合同')
        ck('plan payee 缺省=client', p['context']['payee'] == '测试委托人甲')
        pb = A.api_commission_plan({'fields': {}})
        ck('plan 空委托人拦截（空值纪律）', bool(pb.get('blocked')) and '委托人' in pb['blocked'])
        ps = A.api_commission_plan({'fields': dict(fields, stage='三')})
        ck('plan 非法审级拦截', bool(ps.get('blocked')))
        pa = A.api_commission_plan({'fields': dict(fields, auth_type='全权')})
        ck('plan 非法授权方式拦截', bool(pa.get('blocked')))

        # ---------- build ----------
        rb = A.api_commission_build({'fields': fields, 'case': case}, confirm=False)
        ck('build dry-run 不落盘', rb.get('dry_run')
           and not os.path.isdir(os.path.join(tmp_root, case, '01-委托手续')))
        r1 = A.api_commission_build({'fields': fields, 'case': case}, confirm=True)
        ck('build 六件生成', r1.get('ok') and len(r1['items']) == 6)
        d = os.path.join(tmp_root, case, '01-委托手续')
        ck('落点 01-委托手续', os.path.isdir(d))

        # ---------- 读回 ----------
        d1 = docx.Document(os.path.join(d, '1、委托代理合同（简易版）.docx'))
        t1 = '\n'.join(x.text for x in d1.paragraphs)
        ck('1号 甲方值在填写行', ('测试委托人甲' in t1) and ('甲方（委托方）：' in t1))
        ck('1号 指派/对方/案由/审级', ('安清惠 为甲方' in t1) and ('测试对方乙' in t1)
           and ('买卖合同 纠纷一案' in t1) and ('一 审阶段' in t1))
        ck('1号 律师费', '律师代理费 6000 元' in t1)
        d2 = docx.Document(os.path.join(d, '2、授权委托书-民事.docx'))
        t2 = '\n'.join(x.text for x in d2.paragraphs)
        ck('2号 委托人/代理人与授权方式', ('委托人：测试委托人甲' in t2) and ('律师：安清惠' in t2)
           and ('特别 授权' in t2))
        ck('2号 对方与审级', ('委托人与 测试对方乙' in t2) and ('本人 一 审诉讼代理人' in t2))
        d3 = docx.Document(os.path.join(d, '3、办案反馈表.docx'))
        t3 = '\n'.join(c.text for tb in d3.tables for row in tb.rows for c in row.cells)
        ck('3号 表格填入', all(s in t3 for s in ('测试委托人甲', '民事', '买卖合同纠纷')))
        d5 = docx.Document(os.path.join(d, '5、送达地址确认书-集美法院格式.docx'))
        t5 = '\n'.join(c.text for tb in d5.tables for row in tb.rows for c in row.cells)
        ck('5号 表格填入（当事人/律师）', ('测试委托人甲' in t5) and ('安清惠' in t5) and ('测试地址' in t5))
        d6 = docx.Document(os.path.join(d, '6、账户确认表-集美法院格式.docx'))
        t6 = '\n'.join(c.text for tb in d6.tables for row in tb.rows for c in row.cells)
        ck('6号 表格填入', ('测试委托人甲' in t6) and ('买卖合同纠纷' in t6))

        RE = re.compile(r'<w:t(?: [^>]*)?>(.*?)</w:t>', re.S)
        bad = []
        for f in os.listdir(d):
            with zipfile.ZipFile(os.path.join(d, f)) as z:
                xml = z.read('word/document.xml').decode('utf-8')
            if '{{' in ''.join(x.group(1) for x in RE.finditer(xml)):
                bad.append(f)
        ck('六件无 {{ 残留', not bad)
        allt = t1 + t2 + t3 + t5 + t6
        ck('旧样例数据零残留（消毒）', ('张友谅' not in allt) and ('18950194185' not in allt)
           and ('350211195111083017' not in allt))

        # ---------- 重名避让 ----------
        r2 = A.api_commission_build({'fields': fields, 'case': case}, confirm=True)
        ck('重名避让 -2', any('-2.docx' in x for x in r2['items']))

        # ---------- 特殊字符 ----------
        f2 = dict(fields, client='A&B<C>D公司')
        r3 = A.api_commission_build({'fields': f2, 'case': case}, confirm=True)
        f1 = next(x for x in r3['items'] if x.startswith('1、委托代理合同'))
        d1b = docx.Document(os.path.join(d, f1))
        t1b = '\n'.join(x.text for x in d1b.paragraphs)
        ck('特殊字符不双重转义', ('A&B<C>D公司' in t1b) and ('&amp;' not in t1b))

        # ---------- 安全闸 ----------
        r4 = A.api_commission_build({'fields': fields, 'case': ''}, confirm=True)
        ck('未选案卷拦截', bool(r4.get('error')))
        r5 = A.api_commission_build({'fields': fields, 'case': '../坏夹'}, confirm=True)
        ck('越界夹名拦截', bool(r5.get('error')))
        r6 = A.api_commission_build({'fields': fields, 'case': '不存在的夹'}, confirm=True)
        ck('不存在夹拦截', bool(r6.get('error')))
    finally:
        A.case_ops.case_root = orig_root

    print()
    print('test_api_commission: %d/%d' % (len(total) - len(fails), len(total)))
    return 1 if fails else 0


if __name__ == '__main__':
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:                                        # noqa: BLE001
        pass
    try:
        rc = main()
    finally:
        shutil.rmtree(TMP, ignore_errors=True)
    sys.exit(rc)
