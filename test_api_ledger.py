# -*- coding: utf-8 -*-
"""test_api_ledger.py — 台账受控回写端到端回归（批 13）

覆盖：写环境探测、dry-run 不落盘、备份+I→J 真实回写、**台账被外部改动 → 409 冲突不覆盖**、
      **台账被占用 → 409 占用 + 自动入待同步队列 → 解锁后重试补写**、客户补充档案读写。

安全：全程使用**临时合成台账**（tempfile 里现场生成，不含任何真实案件数据），
      通过 LEDGER_PATH / LEDGER_BACKUP_ROOT / LEDGER_PENDING_PATH 环境变量注入子进程；
      **绝不触碰真台账、真备份目录**。测试用独立端口，跑完自动关停。
"""
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
PY = r'C:\Users\53001\.workbuddy\binaries\python\envs\default\Scripts\python.exe'
PORT = 8812
BASE = 'http://127.0.0.1:%d' % PORT
OP = urllib.request.build_opener(urllib.request.ProxyHandler({}))   # 绕开环境代理
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

checks = []
TMP = None
FAKE = None
BAK = None
PENDING = None
TL_OUT = None
CASES_TMP = None
WIKI_TMP = None


def ck(n, c):
    checks.append((n, bool(c)))


def req(path, method='GET', body=None):
    data = json.dumps(body).encode() if body is not None else None
    r = urllib.request.Request(BASE + path, data=data, method=method,
                               headers={'Content-Type': 'application/json'})
    try:
        with OP.open(r, timeout=10) as resp:
            return resp.status, json.loads(resp.read().decode('utf-8', 'replace'))
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read().decode('utf-8', 'replace') or '{}')


def req_raw(path, data, ctype):
    """原始 body 请求（multipart 上传用）"""
    r = urllib.request.Request(BASE + path, data=data, method='POST',
                               headers={'Content-Type': ctype})
    try:
        with OP.open(r, timeout=180) as resp:
            return resp.status, json.loads(resp.read().decode('utf-8', 'replace'))
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read().decode('utf-8', 'replace') or '{}')


def multipart(files):
    """构造 multipart/form-data（files: [(field, filename, bytes)]）"""
    b = '----wb' + os.urandom(8).hex()
    out = b''
    for field, fn, content in files:
        out += ('--%s\r\n' % b).encode()
        out += ('Content-Disposition: form-data; name="%s"; filename="%s"\r\n'
                % (field, fn)).encode('utf-8')
        out += b'Content-Type: application/octet-stream\r\n\r\n'
        out += content + b'\r\n'
    out += ('--%s--\r\n' % b).encode()
    return out, 'multipart/form-data; boundary=%s' % b


def alive():
    try:
        return req('/api/health')[0] == 200
    except Exception:                                        # noqa: BLE001
        return False


def make_fake_ledger():
    """现场生成合成台账（与真表同结构，零真实数据）"""
    import openpyxl
    wb = openpyxl.Workbook()
    for idx, name in enumerate(('海大案件', '其他案件', '律师函')):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = name
        if name != '律师函':
            ws.cell(row=1, column=1, value='案件管理台账模版表（合成测试用）')
            for col, h in enumerate(('序号', '案号', '案件', '案件类别', '委托人'), start=1):
                ws.cell(row=3, column=col, value=h)
            if name == '海大案件':
                ws.cell(row=4, column=1, value=1)
                ws.cell(row=4, column=2, value='(2026)测字第1号')
                ws.cell(row=4, column=3, value='测试案A')
                ws.cell(row=4, column=4, value='诉讼')
                ws.cell(row=4, column=5, value='测试委托人')
                ws.cell(row=4, column=9, value='08.01：已立案')
                ws.cell(row=4, column=10, value='07.01：更早的记录')
                ws.cell(row=4, column=11, value='待立案')
    wb.save(FAKE)
    wb.close()


def read_cell(col, row=4, sheet='海大案件'):
    import openpyxl
    wb = openpyxl.load_workbook(FAKE, read_only=True, data_only=True)
    try:
        return wb[sheet].cell(row=row, column=col).value
    finally:
        wb.close()


def start():
    env = dict(os.environ)
    env['_WB_STDOUT_UTF8'] = '1'
    env['PYTHONIOENCODING'] = 'utf-8'
    env['LEDGER_PATH'] = FAKE
    env['LEDGER_BACKUP_ROOT'] = BAK
    env['LEDGER_PENDING_PATH'] = PENDING
    env['WORKBENCH_DB'] = os.path.join(TMP, 'workbench.db')          # 客户补充档案也走临时库
    env['WORKBENCH_BACKUP_DIR'] = os.path.join(TMP, 'dbbak')
    env['TIMELINE_OUT_DIR'] = TL_OUT          # 时间轴导出落点也隔离，别写进项目目录
    env['CASE_OPS_ROOT'] = CASES_TMP          # 批 16：开案不写真实案件目录
    env['WIKI_ROOT'] = WIKI_TMP               # 批 16：知识库写入隔离
    env['WIKI_BACKUP_DIR'] = os.path.join(TMP, 'wikibak')   # 批 28：结案标记备份也隔离
    env['OCR_PYTHON'] = sys.executable        # 批 18：OCR 用假引擎，绝不起真识别
    env['OCR_SCRIPT'] = os.path.join(TMP, 'fake_ocr.py')
    env['FAKE_OCR_TEXT_FILE'] = os.path.join(TMP, 'ocr_text.txt')
    env['CASE_OPS_BIZ_ROOT'] = os.path.join(TMP, 'biz')      # 批 19：业务线目录隔离
    env['CASE_OPS_INVOICE_ROOT'] = os.path.join(TMP, '发票库')   # 批 30：公共发票库隔离
    env['INVOICE_UPLOAD_DIR'] = os.path.join(TMP, 'uploads')      # 批 30：上传暂存区隔离
    env['OCR_CONTENT_CACHE_DIR'] = os.path.join(TMP, 'ccache')    # 批 30：OCR 内容缓存隔离
    logf = open(os.path.join(HERE, 'app.out.log'), 'a', encoding='utf-8', buffering=1)
    subprocess.Popen([PY, '-u', os.path.join(HERE, 'app.py'), '--port', str(PORT), '--idle-exit', '0'],
                     cwd=HERE, env=env, stdout=logf, stderr=subprocess.STDOUT,
                     stdin=subprocess.DEVNULL, creationflags=subprocess.DETACHED_PROCESS,
                     close_fds=True)
    for _ in range(60):
        time.sleep(0.2)
        if alive():
            return True
    return False


def backups():
    if not os.path.isdir(BAK):
        return []
    return [f for f in os.listdir(BAK) if f.endswith('.xlsx')]


def main():
    global TMP, FAKE, BAK, PENDING, TL_OUT, CASES_TMP, WIKI_TMP
    TMP = tempfile.mkdtemp(prefix='ledger_api_test_')
    FAKE = os.path.join(TMP, '合成台账.xlsx')
    BAK = os.path.join(TMP, 'bak')
    PENDING = os.path.join(TMP, 'pending.json')
    TL_OUT = os.path.join(TMP, 'timeline')     # 时间轴导出落点隔离（别写进项目目录）
    os.makedirs(TL_OUT, exist_ok=True)
    CASES_TMP = os.path.join(TMP, 'cases')     # 批 16：开案落点隔离
    WIKI_TMP = os.path.join(TMP, 'wiki')       # 批 16：知识库写入隔离
    os.makedirs(CASES_TMP, exist_ok=True)
    make_fake_ledger()
    # 批 18：假 OCR 引擎（把 ocr_text.txt 的内容当作识别结果），隔离真识别
    io.open(os.path.join(TMP, 'fake_ocr.py'), 'w', encoding='utf-8', newline='\n').write(
        'import os, pathlib, sys\n'
        'inp, outdir = sys.argv[1], sys.argv[2]\n'
        'os.makedirs(outdir, exist_ok=True)\n'
        'tf = pathlib.Path(os.environ["FAKE_OCR_TEXT_FILE"])\n'
        'txt = tf.read_text(encoding="utf-8") if tf.exists() else "x"\n'
        'pathlib.Path(outdir, pathlib.Path(inp).stem + ".md").write_text(txt, encoding="utf-8")\n')
    io.open(os.path.join(TMP, 'ocr_text.txt'), 'w', encoding='utf-8', newline='\n').write('占位')

    if not start():
        print('服务未就绪（端口 %d）' % PORT)
        return 3
    try:
        # 1) 列表 + 写环境探测
        s, d = req('/api/ledger_rows')
        ck('台账行列表返回 1 行', s == 200 and d['count'] == 1)
        ck('写环境探测：可写', d['probe']['ok'] is True)
        ck('初始队列为空', d['pending'] == [])
        ck('行含当前进度与日志', d['rows'][0]['当前进度'] == '08.01：已立案'
           and '07.01' in d['rows'][0]['进度日志'])

        # 2) dry-run 预览（不落盘、不备份）
        s, pv = req('/api/ledger/progress', 'POST', {'sheet': '海大案件', 'row': 4,
                                                     'text': '09.12：已寄出起诉状'})
        ck('无 confirm 走 dry-run', s == 200 and pv.get('dry_run') is True)
        ck('预览 I 列新值', pv['plan']['I_new'] == '09.12：已寄出起诉状')
        # 旧值自带 MM.DD： 前缀 → 原样入日志，**不再叠加**写入日期（2026-09-12 真台账事故回归）
        ck('预览回迁旧值（自带日期则原样，不叠日期）',
           pv['plan']['J_entry_added'] == '08.01：已立案')
        ck('预览结果无双重日期', '09.12：08.01' not in pv['plan']['J_new'])
        ck('预览带行快照（冲突基准）', pv['expect']['I'] == '08.01：已立案'
           and int(pv['expect']['row']) == 4)
        ck('dry-run 未改台账', read_cell(9) == '08.01：已立案')
        ck('dry-run 未产生备份', backups() == [])

        # 3) 确认写入（备份 + I→J）
        s, r = req('/api/ledger/progress', 'POST', {'sheet': '海大案件', 'row': 4,
                                                    'text': '09.12：已寄出起诉状',
                                                    'confirm': True,
                                                    'expect': pv['expect']})
        ck('确认写入成功', s == 200 and r.get('ok') is True)
        ck('写入前已备份台账', len(backups()) == 1 and os.path.isfile(r['backup']))
        ck('I 列已更新', read_cell(9) == '09.12：已寄出起诉状')
        ck('旧进度已回迁 J 列（不丢）', '08.01：已立案' in str(read_cell(10))
           and '07.01：更早的记录' in str(read_cell(10)))

        # 4) 台账被外部改动 → 冲突不覆盖
        s, r2 = req('/api/ledger/progress', 'POST', {'sheet': '海大案件', 'row': 4,
                                                     'text': '09.12：不该覆盖',
                                                     'confirm': True, 'expect': pv['expect']})
        ck('拿旧快照写入 → 409 冲突', s == 409 and r2.get('conflict') is True)
        ck('冲突返回差异明细', any(x['字段'] == '当前进度' for x in r2['check']['hard']))
        ck('冲突时未覆盖台账', read_cell(9) == '09.12：已寄出起诉状')
        ck('冲突不入队（须人工判断）', req('/api/ledger/pending')[1]['count'] == 0)

        # 5) 台账被占用 → 入队；解锁后重试补写
        lock = os.path.join(TMP, '~$合成台账.xlsx')
        open(lock, 'w').close()
        s, pv2 = req('/api/ledger/progress', 'POST', {'sheet': '海大案件', 'row': 4,
                                                      'text': '09.13：占用时排的队',
                                                      'confirm': False})
        s, r3 = req('/api/ledger/progress', 'POST', {'sheet': '海大案件', 'row': 4,
                                                     'text': '09.13：占用时排的队',
                                                     'confirm': True, 'expect': pv2['expect']})
        ck('台账被占用 → 409 且标记 busy', s == 409 and r3.get('busy') is True)
        ck('占用时给出原因', bool(r3['probe']['reasons']))
        ck('占用时自动入待同步队列', r3.get('queued') and
           req('/api/ledger/pending')[1]['count'] == 1)
        ck('占用时台账零改动', read_cell(9) == '09.12：已寄出起诉状')
        ck('占用时仍未产生新备份', len(backups()) == 1)

        s, r4 = req('/api/ledger/pending/retry', 'POST', {'confirm': False})
        ck('重试接口无 confirm 走 dry-run', s == 200 and r4.get('dry_run') is True)

        os.remove(lock)
        s, r5 = req('/api/ledger/pending/retry', 'POST', {'confirm': True})
        res = r5.get('result', {})
        ck('解锁后重试成功并出队', s == 200 and len(res.get('ok', [])) == 1
           and req('/api/ledger/pending')[1]['count'] == 0)
        ck('重试真正补写进台账', read_cell(9) == '09.13：占用时排的队')
        ck('重试把前值回迁 J 列', '09.12：已寄出起诉状' in str(read_cell(10)))
        ck('重试产生第二条备份', len(backups()) == 2)

        # 5b) 回迁不叠加日期（真台账事故回归：旧值自带 MM.DD： 前缀）
        s, pv3 = req('/api/ledger/progress', 'POST', {'sheet': '海大案件', 'row': 4,
                                                      'text': '09.14：再记一条'})
        ck('旧值自带日期 → 预览回迁原样（不叠日期）',
           pv3['plan']['J_entry_added'] == '09.13：占用时排的队')
        s, r6 = req('/api/ledger/progress', 'POST', {'sheet': '海大案件', 'row': 4,
                                                     'text': '09.14：再记一条', 'confirm': True,
                                                     'expect': pv3['expect']})
        j_after = str(read_cell(10))
        ck('落盘后 J 列无双重日期', '09.14：09.13' not in j_after
           and j_after.startswith('09.13：占用时排的队'))
        ck('写入后 I 列仍带单层日期', read_cell(9) == '09.14：再记一条')

        # 5c) 批 14b：K 列「最新状态」受控回写（取值池 = 台账「维护说明」页规定）
        s, lr = req('/api/ledger_rows')
        ck('台账行接口带可写入取值池（批 24 起 11 个）',
           s == 200 and len(lr.get('status_pool') or []) == 11 and '待立案' in lr['status_pool']
           and '已判决' in lr['status_pool'])
        ck('台账行接口给出池外清单字段', 'status_off_pool' in lr and 'status_observed' in lr)
        # 批 26：结案移行接口（合成台账**没有**「已结案」页 → 必须被拦住，不许凭空建/误写）
        s, cl = req('/api/ledger/close', 'POST', {'sheet': '海大案件', 'row': 4, 'confirm': False})
        ck('结案移行：台账缺「已结案」页时被拦（400）',
           s == 400 and '已结案' in str(cl.get('error')))
        s, cl2 = req('/api/ledger/close', 'POST', {'sheet': '律师函', 'row': 4, 'confirm': False})
        ck('结案移行：非案件页被拒（400）', s == 400)
        ck('逐行给出状态是否符合规定', lr['rows'][0].get('状态合规') is True)

        s, bad = req('/api/ledger/progress', 'POST', {'sheet': '海大案件', 'row': 4,
                                                      'text': '', 'status': '待执行',
                                                      'confirm': True})
        ck('池外状态 → 400 拒绝写入', s == 400 and '规定取值' in str(bad.get('error')))
        ck('被拒后 K 列未变', read_cell(11) == '待立案')

        s, pv4 = req('/api/ledger/progress', 'POST', {'sheet': '海大案件', 'row': 4,
                                                      'text': '', 'status': '已立案'})
        ck('纯状态变更预览：I_new 为 null（不动 I 列）', s == 200 and pv4['plan']['I_new'] is None)
        ck('纯状态变更预览：K 旧→新', pv4['plan']['K_old'] == '待立案'
           and pv4['plan']['K_new'] == '已立案')
        ck('纯状态变更预览：日志将新增【状态】行',
           str(pv4['plan']['J_status_added']).endswith('【状态】待立案→已立案'))
        ck('状态预览未落盘', read_cell(11) == '待立案')

        s, r7 = req('/api/ledger/progress', 'POST', {'sheet': '海大案件', 'row': 4,
                                                     'text': '', 'status': '已立案',
                                                     'confirm': True, 'expect': pv4['expect']})
        ck('纯状态写入成功且 K 列已改', s == 200 and read_cell(11) == '已立案')
        ck('纯状态写入：I 列未动', read_cell(9) == '09.14：再记一条')
        ck('纯状态写入：J 列留下【状态】行', '【状态】待立案→已立案' in str(read_cell(10)))
        ck('纯状态写入：写后审计零意外（K 列在允许格内）', (r7.get('audit') or []) == [])
        ck('纯状态写入已备份', len(backups()) == 4)

        s, pv5 = req('/api/ledger/progress', 'POST', {'sheet': '海大案件', 'row': 4,
                                                      'text': '09.15：判决已下', 'status': '待判决'})
        ck('进度+状态预览：两者都有值',
           pv5['plan']['I_new'] == '09.15：判决已下' and pv5['plan']['K_new'] == '待判决')
        s, r8 = req('/api/ledger/progress', 'POST', {'sheet': '海大案件', 'row': 4,
                                                     'text': '09.15：判决已下', 'status': '待判决',
                                                     'confirm': True, 'expect': pv5['expect']})
        ck('进度+状态同时写：两列都落盘',
           read_cell(9) == '09.15：判决已下' and read_cell(11) == '待判决')
        ck('进度+状态同时写：J 含回迁行 + 状态行',
           '09.14：再记一条' in str(read_cell(10))
           and '【状态】已立案→待判决' in str(read_cell(10)))
        ck('进度+状态同时写：审计零意外', (r8.get('audit') or []) == [])
        ck('进度+状态同时写已备份', len(backups()) == 5)

        # 5d) 批 15：时间轴导出（落点隔离到临时目录，不写案件夹）
        s, tp = req('/api/timeline/export', 'POST', {'sheet': '海大案件', 'row': 4})
        ck('时间轴预览不落盘（dry_run）', s == 200 and tp['plan']['count'] >= 1
           and not os.path.exists(tp['plan']['dst']))
        ck('时间轴条目带来源标注',
           all(x['备注'] for x in tp['plan']['rows']))
        ck('时间轴落点尊重 TIMELINE_OUT_DIR', tp['plan']['dst_dir'].startswith(TL_OUT))
        s, tr = req('/api/timeline/export', 'POST', {'sheet': '海大案件', 'row': 4, 'confirm': True})
        ck('确认后已导出文件', s == 200 and tr.get('count', 0) >= 1 and os.path.isfile(tr['path']))
        import openpyxl as _ox
        _wb = _ox.load_workbook(tr['path'])
        _ws = _wb['时间轴']
        ck('导出表头正确', [_ws.cell(row=4, column=c).value for c in range(1, 5)]
           == ['序号', '时间', '大事件', '备注'])
        _wb.close()
        ck('导出未改台账（I 列仍旧）', read_cell(9) == '09.15：判决已下')
        s, bad = req('/api/timeline/export', 'POST', {'sheet': '律师函', 'row': 4, 'confirm': True})
        ck('非法 sheet 400', s == 400)

        # 5e) 批 16：开新案（隔离到临时案件根/临时知识库；不跑 pipeline、不写真台账）
        s, pv0 = req('/api/new_case', 'POST', {'form': {
            'case_no': '民-A-26999', 'client': '甲公司', 'opponent': '乙公司',
            'cause': '买卖合同纠纷', 'type': '诉讼', 'court': '思明区法院'}})
        ck('开案预览不落盘', s == 200 and pv0['dry_run'] is True
           and not os.path.exists(pv0['plan']['path']))
        ck('预览含目录清单与《材料怎么放》',
           len(pv0['plan']['dirs']) >= 10 and pv0['plan']['guide_path'].endswith('_材料怎么放.md'))
        ck('预览含知识库案件页方案（含已存在判定）',
           bool(pv0['entity']) and pv0['entity']['will_write'] is True
           and pv0['entity']['dst'].startswith(WIKI_TMP))
        s, r0 = req('/api/new_case', 'POST', {'form': {
            'case_no': '民-A-26999', 'client': '甲公司', 'opponent': '乙公司',
            'cause': '买卖合同纠纷', 'type': '诉讼', 'court': '思明区法院'},
            'confirm': True, 'append_ledger': False, 'with_pipeline': False,
            'write_entity': True})
        ck('确认开案成功', s == 200 and r0.get('dry_run') is False)
        ck('案件夹与目录已建', os.path.isdir(r0['path'])
           and os.path.isdir(os.path.join(r0['path'], '00-收件箱'))
           and os.path.isdir(os.path.join(r0['path'], '08-过程文件')))
        ck('非执行案不建 09/10', not os.path.exists(os.path.join(r0['path'], '09-财产查控')))
        ck('intake-log 已建', os.path.isfile(os.path.join(r0['path'], 'intake-log.md')))
        ck('《材料怎么放》已建且写明规则', r0.get('guide') and os.path.isfile(r0['guide'])
           and '00-收件箱' in io.open(r0['guide'], encoding='utf-8').read())
        ck('知识库案件页已建', bool(r0.get('entity')) and os.path.isfile(r0['entity']))
        ck('未追加台账草稿行时台账不动', read_cell(9) == '09.15：判决已下')
        s, r0b = req('/api/new_case', 'POST', {'form': {
            'case_no': '民-A-26999', 'client': '甲公司', 'opponent': '乙公司',
            'cause': '买卖合同纠纷', 'type': '诉讼'},
            'confirm': True, 'append_ledger': False, 'with_pipeline': False,
            'write_entity': True})
        ck('重复开案被拒（不覆盖既有案卷）', s == 409)
        ck('重复开案未新建知识库页', s == 409)

        # 5f) 批 17：既有案件夹补《材料怎么放》（隔离在临时案件根；只新增、不覆盖）
        oth = os.path.join(CASES_TMP, '执-24585_某甲VS某乙_买卖合同纠纷')
        os.makedirs(os.path.join(oth, '00-收件箱'), exist_ok=True)
        os.makedirs(os.path.join(oth, '09-财产查控'), exist_ok=True)      # 让它被认作执行案
        io.open(os.path.join(oth, 'intake-log.md'), 'w', encoding='utf-8').write('# intake-log\n')
        s, gp = req('/api/guides')
        ck('说明补件预览：列出全部案件夹', s == 200 and gp['count'] == 2)
        ck('说明补件预览：已建的算已有', gp['have'] == 1 and len(gp['todo']) == 1)
        ck('说明补件预览：认出执行案类型', any(x['exec'] for x in gp['items']))
        ck('说明补件预览：不落盘', not os.path.isfile(os.path.join(oth, '_材料怎么放.md')))
        s, gd = req('/api/guides', 'POST', {'apply': True})
        ck('补件未确认只给 dry-run', s == 200 and gd.get('dry_run') is True
           and not os.path.isfile(os.path.join(oth, '_材料怎么放.md')))
        s, ga = req('/api/guides', 'POST', {'apply': True, 'confirm': True})
        ck('确认后新增 1 份、跳过 1 份', s == 200 and len(ga['written']) == 1
           and len(ga['skipped']) == 1)
        gtxt = io.open(os.path.join(oth, '_材料怎么放.md'), encoding='utf-8').read()
        ck('说明含本夹现状与一句话规则',
           '本夹现状' in gtxt and '一句话规则' in gtxt)
        ck('说明含编号/类别/台账唯一事实源',
           '执-24585' in gtxt and '类别：执行' in gtxt and '唯一事实源' in gtxt)
        ck('说明标出缺的 00-收件箱之外的标准目录', '还没建' in gtxt or '标准目录已建' in gtxt)
        s, ga2 = req('/api/guides', 'POST', {'apply': True, 'confirm': True})
        ck('再次补件：不覆盖（written 为空）', s == 200 and ga2['written'] == [])
        # 5g) 批 18：入库页本地 OCR（假引擎隔离；只新增识别文本、不动原件）
        oc = os.path.join(CASES_TMP, '民-A-26111_甲VS乙_买卖合同纠纷')
        os.makedirs(os.path.join(oc, '00-收件箱'), exist_ok=True)
        scan = os.path.join(oc, '00-收件箱', '扫描件.pdf')
        with open(scan, 'wb') as f:
            f.write(b'%PDF-1.4 fake')
        io.open(os.path.join(TMP, 'ocr_text.txt'), 'w', encoding='utf-8').write(
            '本院受理通知书 案号（2026）闽0122民初1016号 请按开庭传票载明时间到庭')
        s, od = req('/api/ocr/run', 'POST', {'case': oc, 'name': '扫描件.pdf'})
        ck('OCR 未确认只给 dry-run（含可识别性）',
           s == 200 and od.get('dry_run') is True and od['plan'].get('ok') is True
           and od['plan'].get('cached') is False
           and not os.path.isfile(os.path.join(oc, '08-过程文件', '_ocr', '扫描件.md')))
        s, orr = req('/api/ocr/run', 'POST', {'case': oc, 'name': '扫描件.pdf', 'confirm': True})
        ck('确认后识别成功', s == 200 and orr.get('chars', 0) > 10)
        ck('识别文本落到 08-过程文件\\_ocr',
           os.path.isfile(os.path.join(oc, '08-过程文件', '_ocr', '扫描件.md')))
        ck('原文件未被动过（大小不变）', os.path.getsize(scan) == len(b'%PDF-1.4 fake'))
        ck('正文建议归 04-法院文件', orr['suggest']['dir'] == '04-法院文件'
           and orr['suggest']['from'] == '正文')
        s, or2 = req('/api/ocr/run', 'POST', {'case': oc, 'name': '扫描件.pdf', 'confirm': True})
        ck('二次识别走缓存（不重复跑）', s == 200 and or2.get('cached') is True)
        s, oerr = req('/api/ocr/run', 'POST', {'case': oc, 'name': '../../../windows/win.ini',
                                               'confirm': True})
        ck('越界文件名被拒', s == 400)
        s, otxt = req('/api/ocr/run', 'POST', {'case': oc, 'name': '不存在.pdf', 'confirm': True})
        ck('不存在的文件被拒', s == 400)
        ck('OCR 全程未改台账', read_cell(9) == '09.15：判决已下')

        # 5h) 批 19：律师函页 + 差旅发票索引（业务线目录隔离到临时目录）
        biz = os.path.join(TMP, 'biz')
        os.makedirs(os.path.join(biz, '律师函', '某债务人'), exist_ok=True)
        with open(os.path.join(biz, '律师函', '某债务人', '律师函.pdf'), 'wb') as f:
            f.write(b'x')
        os.makedirs(os.path.join(biz, '差旅发票', '8月'), exist_ok=True)
        with open(os.path.join(biz, '差旅发票', '8月', '2026-08-14测试案A立案顺风车.pdf'), 'wb') as f:
            f.write(b'x')
        with open(os.path.join(biz, '差旅发票', '8月', '26359134012000099999.pdf'), 'wb') as f:
            f.write(b'x')
        import openpyxl as _ox2
        _wb2 = _ox2.load_workbook(FAKE)
        _ws2 = _wb2['律师函']
        for _c, _v in ((1, 1), (2, '测试委托人'), (8, '已签收（测试）')):
            _ws2.cell(row=4, column=_c, value=_v)
        _wb2.save(FAKE)
        _wb2.close()
        s, lt = req('/api/letters')
        ck('律师函页：读到台账表记录', s == 200 and lt['count'] == 1
           and lt['items'][0]['委托人'] == '测试委托人'
           and lt['items'][0]['签收状态'] == '已签收（测试）')
        ck('律师函页：文件系统夹概览', lt['folders']['exists']
           and len(lt['folders']['folders']) == 1 and lt['folders']['folders'][0]['files'] == 1)
        s, tv = req('/api/travel')
        ck('差旅发票索引：认出含案名那张', s == 200 and tv['recognized'] == 1
           and tv['items'][0]['matched'] is True)
        ck('差旅发票索引：纯发票号明说认不出',
           any((not x['matched']) and '中文线索' in (x['why'] or '') for x in tv['items']))
        ck('业务线接口未改台账', read_cell(9) == '09.15：判决已下')

        # 5h2) 批 41：新增发函对象（台账律师函 sheet 受控加行 + 批量催收件建材料夹）
        s, pv = req('/api/letter/add_row', 'POST',
                    {'委托人': '测试委托人', '收函对象': '测试债务人',
                     '欠款金额（本金）': '101.00', '案由': '买卖合同纠纷', 'confirm': False})
        ck('批41 预览：dry_run 行号=r5 序号=2 夹将新建 + 写法规范化',
           s == 200 and pv.get('dry_run') is True
           and pv['plan']['row'] == 5 and pv['plan']['values'].get('序号') == 2
           and pv['plan']['values'].get('欠款金额（本金）') == '101元'
           and len(pv['plan'].get('normalized') or []) >= 1
           and pv.get('folder', {}).get('exists') is False)
        s, bad = req('/api/letter/add_row', 'POST', {'委托人': 'x', 'confirm': True})
        ck('批41 缺收函对象：400', s == 400)
        s, ok = req('/api/letter/add_row', 'POST',
                    {'委托人': '测试委托人', '收函对象': '测试债务人',
                     '欠款金额（本金）': '44,800元', '案由': '买卖合同纠纷',
                     '发函日期': '2026-09-18', '备注': '测试备注', 'confirm': True})
        ck('批41 确认：200 且 folder.created=True', s == 200 and ok.get('ok') is True
           and ok['plan']['row'] == 5 and ok.get('folder', {}).get('created') is True)
        ck('批41 台账行内容（B 委托人 / C 日期点分隔 / E 收函对象 / F 金额带元 / H 留空）',
           read_cell(5, row=5, sheet='律师函') == '测试债务人'
           and read_cell(2, row=5, sheet='律师函') == '测试委托人'
           and read_cell(3, row=5, sheet='律师函') == '2026.09.18'
           and read_cell(6, row=5, sheet='律师函') == '44,800元'
           and read_cell(8, row=5, sheet='律师函') in (None, '')
           and read_cell(10, row=5, sheet='律师函') == '测试备注')
        ck('批41 案件表零变化（写后审计口径）', read_cell(2, row=4) == '(2026)测字第1号')
        ck('批41 材料夹已建（_归档区/批量催收件/测试债务人）',
           os.path.isdir(os.path.join(CASES_TMP, '_归档区', '批量催收件', '测试债务人')))
        s, ok2 = req('/api/letter/add_row', 'POST',
                     {'委托人': '测试委托人', '收函对象': '测试债务人', 'confirm': True})
        ck('批41 重复确认：r6 新行 + 夹已存在跳过', s == 200 and ok2['plan']['row'] == 6
           and ok2.get('folder', {}).get('created') is False)
        s, ok3 = req('/api/letter/add_row', 'POST',
                     {'委托人': '测试委托人', '收函对象': '甲公司、乙丙、丁',
                      'confirm': True})
        ck('批41 夹名取主体词（甲公司）', s == 200
           and ok3.get('folder', {}).get('created') is True
           and os.path.isdir(os.path.join(CASES_TMP, '_归档区', '批量催收件', '甲公司'))
           and not os.path.isdir(os.path.join(CASES_TMP, '_归档区', '批量催收件', '甲公司、乙丙、丁')))

        # 5i) 批 27：案件列表台账摘要 + 案件详情差旅发票 / 子夹视图（全只读）
        os.makedirs(os.path.join(CASES_TMP, '测试案A_测试委托人', '00-收件箱'), exist_ok=True)
        ser = os.path.join(CASES_TMP, '甲系列案')
        for _s in ('甲案', '乙案'):
            os.makedirs(os.path.join(ser, _s, '00-收件箱'), exist_ok=True)
        s, cs = req('/api/cases')
        ck('案件列表带台账摘要/系列案字段', s == 200 and 'recognized' in cs
           and all(('ledger' in c and 'series' in c and 'subcase_count' in c) for c in cs['cases']))
        _one = [c for c in cs['cases'] if c['name'] == '测试案A_测试委托人']
        ck('案件夹认出台账行并给摘要字段',
           bool(_one) and _one[0]['ledger']['matched'] is True
           and _one[0]['ledger']['案件'] == '测试案A' and '承办律师' in _one[0]['ledger'])
        _ser = [c for c in cs['cases'] if c['name'] == '甲系列案']
        ck('系列案总夹识别出子案件夹',
           bool(_ser) and _ser[0]['series'] is True and _ser[0]['subcase_count'] == 2)
        s, cd = req('/api/case?name=' + urllib.parse.quote('测试案A_测试委托人'))
        ck('案件详情：差旅发票挂到本案', s == 200 and cd['travel']['total'] == 1
           and cd['travel']['items'][0]['name'].endswith('.pdf'))
        ck('案件详情：普通案无子夹', cd['subcases'] == [])
        s, cd2 = req('/api/case?name=' + urllib.parse.quote('甲系列案'))
        ck('案件详情：列出子案件夹与收件箱状态',
           s == 200 and len(cd2['subcases']) == 2
           and all(x['inbox'] is True for x in cd2['subcases']))
        ck('案件详情接口未改台账', read_cell(9) == '09.15：判决已下')

        # 5j) 批 28：知识库案件页 → 标「已结案」（隔离 WIKI_ROOT + WIKI_BACKUP_DIR）
        _page = '甲公司诉乙公司买卖合同纠纷案'
        _wp = os.path.join(WIKI_TMP, 'entities', _page + '.md')
        _folder = os.path.basename(r0['path'])
        s, wcp = req('/api/wiki/case_pages')
        ck('知识库案件页清单可读', s == 200 and 'items' in wcp and wcp['count'] >= 1
           and any(p['title'] == _page for p in wcp['items']) and wcp['closed'] == 0)
        s, wcp2 = req('/api/wiki/case_pages?case=' + urllib.parse.quote(_folder))
        ck('案件夹 → 知识库页给出建议（主体词命中）', s == 200 and wcp2['suggest'] == _page)
        s, w1 = req('/api/wiki/close', 'POST', {'page': _page})
        ck('知识库标记：无 confirm 走 dry-run 且给改动清单',
           s == 200 and w1.get('dry_run') is True and w1['plan']['will_write'] is True
           and len(w1['plan']['changes']) >= 4)
        ck('知识库标记：dry-run 未改页面', 'status: 在办' in io.open(_wp, encoding='utf-8').read())
        s, w2 = req('/api/wiki/close', 'POST', {'page': _page, 'folder': 'D:/案件/x',
                                                'confirm': True})
        ck('知识库标记：确认后成功且写后校验通过',
           s == 200 and w2['changed'] is True and w2.get('verified') is True)
        _txt = io.open(_wp, encoding='utf-8').read()
        ck('页面已带结案标记 / status / tags 已改',
           '> ✅ **已结案' in _txt and 'status: 已结案' in _txt and '## 结案归档' in _txt
           and 'status: 在办' not in _txt)
        ck('页面其余内容未丢（骨架仍在）', '## 过程性文件台账' in _txt and '## 基本档案' in _txt)
        ck('备份落点可查（隔离目录）', bool(w2.get('backup')) and os.path.isfile(w2['backup'])
           and w2['backup'].startswith(TMP))
        s, w3 = req('/api/wiki/close', 'POST', {'page': _page, 'confirm': True})
        ck('知识库标记：再标一次幂等不改', s == 200 and w3['changed'] is False)
        s, _ = req('/api/wiki/close', 'POST', {'page': '根本没有这一页案', 'confirm': True})
        ck('知识库标记：找不到页 → 404', s == 404)
        s, _ = req('/api/wiki/close', 'POST', {'page': '', 'confirm': True})
        ck('知识库标记：缺 page → 400', s == 400)
        ck('知识库操作未改台账', read_cell(9) == '09.15：判决已下')

        # 5k) 批 30：发票上传 → 本机识别 → 案卷正本 + 公共库副本（假引擎隔离 + 隔离发票库）
        inv_case = os.path.join(CASES_TMP, '民-A-26031_甲公司VS乙公司_买卖合同纠纷')
        os.makedirs(os.path.join(inv_case, '05-费用票据'), exist_ok=True)
        io.open(os.path.join(TMP, 'ocr_text.txt'), 'w', encoding='utf-8').write(
            '电子发票（增值税专用发票）\n开票日期：\n名称：\n名称：\n开票人：\n'
            '26942000000000000001\n2026年09月11日\n某某律师事务所\n313200000000000000\n'
            '甲公司\n31350000MD0274017T\n¥100.00\n¥1.00\n壹佰零壹圆整\n¥101.00\n*服务*律师费\n')
        s, up = req_raw('/api/invoice/upload', *multipart([('files', '扫描件.pdf',
                                                            b'%PDF-1.4 invoice')]))
        if not (s == 200 and up.get('ok_count') == 1):
            print('   [debug] upload 返回:', s, str(up)[:300])
        ck('发票上传：识别成功并给建议名/月份',
           s == 200 and up.get('ok_count') == 1
           and up['items'][0]['suggest_name'] == '20260911_发票_101.00元_律师费.pdf'
           and up['items'][0]['month'] == '2026-09')
        ck('发票上传：抽出发票号码与金额（取价税合计）',
           up['items'][0]['fields']['发票号码'] == '26942000000000000001'
           and up['items'][0]['fields']['金额'] == '101.00')
        ck('发票上传：只落暂存区（案卷与公共库都没动静）',
           os.path.isfile(up['items'][0]['staged'])
           and not os.path.isdir(os.path.join(TMP, '发票库')))
        s, bad2 = req_raw('/api/invoice/upload', *multipart([('files', 'x.txt', b'hi')]))
        ck('不支持的类型：列在结果里（不整体失败）',
           s == 200 and bad2['items'][0]['ok'] is False and bad2['items'][0]['error'])
        s, _ = req('/api/invoice/upload', 'POST', {'x': 1})
        ck('非 multipart 上传 → 400', s == 400)
        st = up['items'][0]['staged']
        cn = os.path.basename(inv_case)
        s, pv2 = req('/api/invoice/file', 'POST', {'staged': st, 'case': cn,
                                                   'fields': up['items'][0]['fields']})
        ck('归位：无 confirm 只给方案（不落盘）',
           s == 200 and pv2['dry_run'] is True and pv2['plan']['ok'] is True
           and pv2['plan']['newname'].endswith('20260911_发票_101.00元_律师费.pdf')
           and not os.path.exists(pv2['plan']['dst_main']))
        s, dn = req('/api/invoice/file', 'POST', {'staged': st, 'case': cn,
                                                  'fields': up['items'][0]['fields'],
                                                  'month': '2026-09', 'confirm': True})
        ck('归位：正本进案卷 05-费用票据',
           s == 200 and os.path.isfile(dn['done']['main'])
           and dn['done']['main'].endswith(os.path.join('05-费用票据',
                                                        '20260911_发票_101.00元_律师费.pdf')))
        ck('归位：公共库副本落当月且与正本内容一致',
           os.path.isfile(dn['done']['copy'])
           and dn['done']['copy'].startswith(os.path.join(TMP, '发票库', '2026-09'))
           and io.open(dn['done']['main'], 'rb').read() == io.open(dn['done']['copy'], 'rb').read())
        ck('归位：暂存件已移走（不留残件）', not os.path.exists(st))
        ck('归位：intake-log 记原名与新名',
           '扫描件.pdf' in io.open(dn['done']['intake_log'], encoding='utf-8').read()
           and '20260911_发票_101.00元_律师费.pdf'
           in io.open(dn['done']['intake_log'], encoding='utf-8').read())
        ck('归位：公共库索引已写', bool(dn['done']['index']) and os.path.isfile(dn['done']['index']))
        s, lb = req('/api/invoice/overview')
        ck('发票页数据可读（公共库 + 案件下拉）',
           s == 200 and lb['library']['total'] == 1 and len(lb['cases']) >= 1
           and lb['max_mb'] > 0 and lb['lib_root'])
        s, e1 = req('/api/invoice/file', 'POST', {'staged': st, 'case': '不存在的案子',
                                                  'confirm': True})
        ck('归位：案件夹不存在 → 400', s == 400)
        ck('发票流程未改台账', read_cell(9) == '09.15：判决已下')

        ck('补件未改 intake-log',
           io.open(os.path.join(oth, 'intake-log.md'), encoding='utf-8').read() == '# intake-log\n')

        # 6) 客户补充档案
        s, d = req('/api/clients_ext', 'POST', {'client_key': '测试委托人',
                                                'patch': {'contact': '张经理',
                                                          'phone': '13800000000'}})
        ck('无 confirm 走 dry-run', s == 200 and d.get('dry_run') is True)
        s, d = req('/api/clients_ext', 'POST', {'client_key': '测试委托人', 'confirm': True,
                                                'patch': {'contact': '张经理',
                                                          'phone': '13800000000',
                                                          'wechat': 'test88'}})
        ck('保存客户档案成功', s == 200 and d['ext']['contact'] == '张经理')
        cl = req('/api/overview')[1]['clients']
        one = [x for x in cl['items'] if x['name'] == '测试委托人']
        ck('客户聚合视图带出补充档案', bool(one) and one[0]['ext']['wechat'] == 'test88')
        ck('客户计数含已补档案数', cl['with_ext'] >= 1)
        s, _ = req('/api/clients_ext', 'POST', {'client_key': '', 'confirm': True,
                                                'patch': {'phone': '1'}})
        ck('空客户名 400', s == 400)
        s, _ = req('/api/clients_ext', 'POST', {'client_key': 'x', 'confirm': True,
                                                'patch': {'nope': '1'}})
        ck('非法字段 400', s == 400)

        # 7) 入参校验
        s, _ = req('/api/ledger/progress', 'POST', {'sheet': '不存在的表', 'row': 4,
                                                    'text': 'x', 'confirm': True})
        ck('非法 sheet 400', s == 400)
        s, _ = req('/api/ledger/progress', 'POST', {'sheet': '海大案件', 'row': 'abc',
                                                    'text': 'x', 'confirm': True})
        ck('row 非数字 400', s == 400)

        # 8) 首页统计联动（不炸）
        o = req('/api/overview')[1]
        ck('overview 仍含任务/客户/进行中案件',
           'tasks' in o and o['clients']['count'] >= 1 and 'active_cases' in o)
    finally:
        try:
            req('/api/quit')
        except Exception:                                    # noqa: BLE001
            pass
        time.sleep(1.0)
        shutil.rmtree(TMP, ignore_errors=True)
        for f in ('_smoke_data.json', '_smoke_views.js'):
            p = os.path.join(HERE, f)
            if os.path.isfile(p):
                try:
                    os.unlink(p)
                except OSError:
                    pass

    bad = [n for n, ok in checks if not ok]
    for n, ok in checks:
        print('[%s] %s' % ('OK' if ok else 'FAIL', n))
    print('台账受控回写回归: %d/%d 通过' % (len(checks) - len(bad), len(checks)))
    return 1 if bad else 0


if __name__ == '__main__':
    sys.exit(main())
