# -*- coding: utf-8 -*-
"""批 39b 隔离回归：律师函表受控回写（**全程合成临时台账，不碰真数据**）。

覆盖：计划校验 / 拒绝空改 / 同值跳过 / 行号越界 / 备份生成 / 写后复查 /
      J 列（备注）绝不被动 / 冲突检测 / 只允许两列。
"""
import io
import os
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)

TMP = tempfile.mkdtemp(prefix='p39b_ledger_')
LEDGER = os.path.join(TMP, '合成台账.xlsx')
BK = os.path.join(TMP, 'backup')
os.environ['LEDGER_PATH'] = LEDGER
os.environ['LEDGER_BACKUP_ROOT'] = BK

fails, total = [], []


def ck(name, cond):
    print('%s %s' % ('PASS' if cond else 'FAIL', name))
    total.append(name)
    if not cond:
        fails.append(name)


def _raises(fn):
    try:
        fn()
        return False
    except Exception:                                        # noqa: BLE001
        return True


# ---- 合成台账：律师函表 A-J，表头 3 行，数据自 r4 ----
import openpyxl                                           # noqa: E402


def build():
    wb = openpyxl.Workbook()
    for nm in ('海大案件', '其他案件', '律师函'):
        wb.create_sheet(nm)
    wb.remove(wb['Sheet'])
    ws = wb['律师函']
    for c, v in enumerate(['序号', '委托人', '发函日期', '承办律师', '收函对象',
                           '欠款金额（本金）', '案由', '签收状态', '后续跟进', '备注'], 1):
        ws.cell(row=3, column=c, value=v)
    ws.cell(row=4, column=1, value='1')
    ws.cell(row=4, column=2, value='甲公司')
    ws.cell(row=4, column=5, value='乙、丙')
    ws.cell(row=4, column=8, value='未签收')
    ws.cell(row=4, column=9, value='')
    ws.cell(row=4, column=10, value='备注原文-勿动')
    ws.cell(row=5, column=1, value='2')
    ws.cell(row=5, column=2, value='丁公司')
    ws.cell(row=5, column=5, value='戊')
    ws.cell(row=5, column=8, value='已签收')
    ws.cell(row=5, column=9, value='已电话联系')
    ws.cell(row=5, column=10, value='另一备注')
    wb.save(LEDGER)
    wb.close()


build()
import ledger_ops                                        # noqa: E402

# ① 计划：改两列
p = ledger_ops.plan_letter_update(4, sign='3封已签收', follow='已电话联系，承诺本周还')
ck('计划拿到 2 处改动', len(p['changes']) == 2)
ck('计划含现状快照', p['现状'].get('委托人') == '甲公司')

# ② 拒绝空改
ck('两列都空被拒', _raises(lambda: ledger_ops.plan_letter_update(4)))
ck('与现状同值被拒', _raises(lambda: ledger_ops.plan_letter_update(5, sign='已签收')))
# ③ 行号越界
ck('row<4 被拒', _raises(lambda: ledger_ops.plan_letter_update(2, sign='x')))
ck('不存在的行被拒', _raises(lambda: ledger_ops.plan_letter_update(999, sign='x')))
# ④ 只允许两列
ck('可写列只有签收状态/后续跟进',
   set(ledger_ops.LETTER_WRITE_COLS.keys()) == {'签收状态', '后续跟进'})
ck('备注列不在可写列内', '备注' not in ledger_ops.LETTER_WRITE_COLS)

# ⑤ 执行（带 expect 正常）
r = ledger_ops.apply_letter_update(4, sign='3封已签收', follow='已电话联系，承诺本周还',
                                   expect={'签收状态': '未签收', '后续跟进': ''})
ck('执行成功', r.get('ok') is True)
ck('备份已生成', bool(r.get('backup')) and os.path.isfile(r['backup']))
after = [x for x in ledger_ops.read_rows('律师函') if x['_row'] == 4][0]
ck('签收状态已写', str(after['签收状态']).strip() == '3封已签收')
ck('后续跟进已写', str(after['后续跟进']).strip() == '已电话联系，承诺本周还')
ck('备注列未被改', str(after['备注']).strip() == '备注原文-勿动')
ck('写后审计零意外', not (r.get('audit') or []))

# ⑥ 冲突检测（expect 与现状不符 → 409 语义）
r2 = ledger_ops.apply_letter_update(4, sign='乱写', expect={'签收状态': '过时的值'})
ck('expect 不符 → conflict', r2.get('conflict') is True)
ck('冲突时不写真数据', str([x for x in ledger_ops.read_rows('律师函')
                        if x['_row'] == 4][0]['签收状态']).strip() == '3封已签收')

# ⑦ 只传一列（另一列不动）
r3 = ledger_ops.apply_letter_update(5, follow='新跟进内容',
                                    expect={'签收状态': '已签收', '后续跟进': '已电话联系'})
ck('单列写入成功', r3.get('ok') is True)
a5 = [x for x in ledger_ops.read_rows('律师函') if x['_row'] == 5][0]
ck('只改了后续跟进', str(a5['后续跟进']).strip() == '新跟进内容')
ck('签收状态保持原值', str(a5['签收状态']).strip() == '已签收')

# ⑧ 备份数量
bks = [f for f in os.listdir(BK) if f.endswith('.xlsx')]
ck('备份至少 2 份（两次写入各一份）', len(bks) >= 2)

shutil.rmtree(TMP, ignore_errors=True)
print('--- %d 项，失败 %d ---' % (len(total), len(fails)))
if fails:
    print('FAILED: %s' % fails)
sys.exit(1 if fails else 0)
