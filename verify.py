#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""verify.py — 一键自检：clone 下来先跑这个，比看文档快

检查项：
  1. Python 版本与必需/可选依赖
  2. 配置文件来源与关键路径
  3. 台账可读、sheet 与列映射对得上
  4. 案卷根 / 收件箱结构
  5. 端口占用
  6. 演示数据是否就绪

用法：
    python verify.py
退出码：0 可用 / 1 有阻断项
"""
import importlib
import os
import socket
import sys

import config

OK, WARN, BAD = [], [], []


def ok(msg):
    OK.append(msg)


def warn(msg):
    WARN.append(msg)


def bad(msg):
    BAD.append(msg)


def check_python():
    v = sys.version_info
    if v >= (3, 10):
        ok('Python %d.%d.%d' % v[:3])
    else:
        bad('Python %d.%d 过低（需要 3.10+）' % v[:2])
    for mod, need in (('openpyxl', True), ('fitz', False), ('docxtpl', False),
                      ('webview', False), ('docx', False)):
        try:
            importlib.import_module(mod)
            ok('依赖 %s 已安装' % mod)
        except ImportError:
            (bad if need else warn)('依赖 %s 未安装%s' % (mod, '（必需）' if need else '（可选，对应功能不可用）'))


def check_config():
    m = config.CFG.get('_meta') or {}
    if m.get('local_exists'):
        ok('配置：使用本机 workbench.config.json')
    else:
        warn('配置：未发现 workbench.config.json，正在用示例配置（demo 数据可直接跑）')
    demo_dir = os.path.join(config.BASE_DIR, 'demo')
    demo_ready = os.path.isdir(demo_dir)
    for key in ('ledger', 'case_root', 'archive_root', 'ledger_backup_root'):
        p = config.abspath(config.get('paths.' + key, ''))
        if not p:
            bad('配置 paths.%s 为空' % key)
        elif os.path.exists(p):
            ok('paths.%s 存在：%s' % (key, p))
        elif ('demo' in p.split(os.sep)) and not demo_ready:
            # 首次使用还没生成示例数据：这不是错误，别让新人和 agent 以为装坏了
            warn('paths.%s 待生成（首次使用正常现象）：%s' % (key, p))
        else:
            warn('paths.%s 不存在：%s' % (key, p))


def check_ledger():
    p = config.abspath(config.get('paths.ledger', ''))
    if not os.path.isfile(p):
        if 'demo' in p.split(os.sep) and not os.path.isdir(
                os.path.join(config.BASE_DIR, 'demo')):
            warn('台账待生成（首次使用正常现象）：%s' % p)
        else:
            bad('台账文件不存在：%s' % p)
        return
    try:
        from openpyxl import load_workbook
    except ImportError:
        return
    try:
        wb = load_workbook(p, read_only=True, data_only=True)
    except Exception as e:                                   # noqa: BLE001
        bad('台账打不开（可能被 WPS/Excel 占用）：%r' % (e,))
        return
    try:
        sheets = config.get('ledger.sheets_case') or []
        missing = [s for s in sheets if s not in wb.sheetnames]
        if missing:
            bad('台账缺 sheet：%s（现有：%s）' % ('/'.join(missing), '/'.join(wb.sheetnames)))
        else:
            ok('台账 sheet 齐全：%s' % '/'.join(sheets))
        ws = wb[sheets[0]] if sheets and sheets[0] in wb.sheetnames else None
        if ws is not None:
            cols = config.get('ledger.cols_case') or {}
            hdr_row = int(config.get('ledger.data_start_row', 4)) - 1
            names = [str(c.value or '').strip() for c in next(
                ws.iter_rows(min_row=hdr_row, max_row=hdr_row))]
            hit = sum(1 for k in cols if k in names)
            if hit >= max(3, len(cols) // 2):
                ok('列映射对得上（表头 %d 列中命中 %d 个字段名）' % (len(names), hit))
            else:
                warn('列映射可能不对：表头第 %d 行只命中 %d 个字段名，请核对 ledger.cols_case'
                     % (hdr_row, hit))
            must = ('当前进度', '进度日志', '最新状态')
            miss = [k for k in must if k not in cols or k not in names]
            if miss:
                bad('必填列缺失或表头对不上：%s（受控回写依赖这三列）' % '/'.join(miss))
            else:
                ok('三个必填列就位：%s' % '/'.join(must))
    finally:
        wb.close()


def check_cases():
    root = config.abspath(config.get('paths.case_root', ''))
    if not os.path.isdir(root):
        if 'demo' in root.split(os.sep) and not os.path.isdir(
                os.path.join(config.BASE_DIR, 'demo')):
            warn('案卷根待生成（首次使用正常现象）：%s' % root)
        else:
            bad('案卷根不存在：%s' % root)
        return
    subs = [d for d in os.listdir(root) if os.path.isdir(os.path.join(root, d))]
    if not subs:
        warn('案卷根下没有案件夹：%s' % root)
        return
    with_inbox = [d for d in subs if os.path.isdir(os.path.join(root, d, '00-收件箱'))]
    ok('案卷 %d 个，其中含 00-收件箱 的 %d 个' % (len(subs), len(with_inbox)))


def check_port():
    port = int(config.get('app.port', 8899) or 8899)
    s = socket.socket()
    s.settimeout(0.3)
    try:
        s.connect(('127.0.0.1', port))
        warn('端口 %d 已被占用（可能是工作台已在运行）' % port)
    except OSError:
        ok('端口 %d 可用' % port)
    finally:
        s.close()


def main():
    print('=== verify: 法律工作台自检 ===\n')
    check_python()
    check_config()
    check_ledger()
    check_cases()
    check_port()

    for m in OK:
        print('  ✅ ' + m)
    for m in WARN:
        print('  ⚠️  ' + m)
    for m in BAD:
        print('  ❌ ' + m)

    print('\n结果：%d 项通过 / %d 项提醒 / %d 项阻断' % (len(OK), len(WARN), len(BAD)))
    if BAD:
        print('\n下一步：')
        print('  python demo_seed.py     # 生成示例数据')
        print('  python config.py        # 看当前生效的路径')
        return 1
    if any('待生成' in m for m in WARN):
        print('\n下一步： python demo_seed.py   # 生成示例数据后即可启动')
    print('\n可以启动了： python app.py')
    return 0


if __name__ == '__main__':
    sys.exit(main())
