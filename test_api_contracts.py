# -*- coding: utf-8 -*-
r"""test_api_contracts.py — 合同审查业务线回归（批 52，2026-10-09）

覆盖：
  · 行解析（序号/合同名称/收费原样字符）
  · 收费合计（H 列求和；非数字跳过）；状态计数
  · sheet 缺失 → ok + count 0（read_rows guard 优雅返回 []）
  · 素材夹缺失 → folder.exists False；夹存在 → files 列表（kind 字段）
  · 全页只读：调用前后台账文件 mtime 不变

⚠️ 全程合成临时台账（openpyxl 现场生成）+ 临时素材根（CASE_OPS_BIZ_ROOT），
   **不碰真台账、不碰真素材夹**。
"""
import os
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

TMP = tempfile.mkdtemp(prefix='wb_contracts_test_')
os.environ['CASE_OPS_BIZ_ROOT'] = os.path.join(TMP, 'biz')

fails, total = [], []


def ck(name, cond):
    total.append(name)
    print(('  PASS ' if cond else '  FAIL ') + name)
    if not cond:
        fails.append(name)


def _make_ledger(path, with_contract=True):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '海大案件'
    ws.cell(row=1, column=1, value='测试')
    if with_contract:
        wc = wb.create_sheet('合同审查')
        for i, h in enumerate(['序号', '送审日期', '委托人', '合同名称', '相对方',
                               '合同类型', '审查状态', '收费', '交付日期', '备注'], start=1):
            wc.cell(row=3, column=i, value=h)
        rows = [
            (1, '2026-09-20', '甲公司', '测试合同A（买卖合同）', '乙公司', '买卖', '待审', '5,000', '', ''),
            (2, '2026-09-25', '丙公司', '测试合同B（服务合同）', '丁公司', '服务', '审查中', 3500.5, '', ''),
            (3, '2026-10-01', '戊公司', '测试合同C（租赁合同）', '己公司', '租赁', '已定稿', 'abc', '2026-10-05', ''),
        ]
        for ri, r in enumerate(rows, start=4):
            for ci, v in enumerate(r, start=1):
                wc.cell(row=ri, column=ci, value=v)
        # 第 7 行留空（空行跳过断言）
    wb.save(path)
    wb.close()


def main():
    import app as A
    import ledger_ops as L

    ledger = os.path.join(TMP, '台账.xlsx')
    _make_ledger(ledger, with_contract=True)
    os.environ['LEDGER_PATH'] = ledger

    print('=== ① sheet 存在：解析/合计/计数 ===')
    mtime0 = os.path.getmtime(ledger)
    r = A.api_contracts()
    ck('ok + count=3（空行跳过）', r.get('ok') is True and r['count'] == 3)
    it = r['items'][0]
    ck('行解析（序号/合同名称/收费原样）',
       it['序号'] == '1' and it['合同名称'] == '测试合同A（买卖合同）' and it['收费'] == '5,000')
    ck('row 字段=台账真实行号', it['row'] == 4)
    ck('第 2 行数字收费原样保留', r['items'][1]['收费'] == '3500.5')
    ck('收费合计 = 5000 + 3500.5（"abc" 跳过）',
       r['fee_total'] == 8500.5 and r['fee_count'] == 2)
    ck('状态计数', r['status_counts'] == {'待审': 1, '审查中': 1, '已定稿': 1})
    ck('sheet 名回给前端', r['sheet'] == '合同审查')

    print('=== ② 素材夹 ===')
    ck('夹缺失 → exists False + 落点在 biz_root\\合同审查',
       r['folder']['exists'] is False
       and r['folder']['root'] == os.path.join(TMP, 'biz', '合同审查'))
    fdir = os.path.join(TMP, 'biz', '合同审查')
    os.makedirs(fdir)
    with open(os.path.join(fdir, '待审合同-甲公司.pdf'), 'wb') as f:
        f.write(b'%PDF-1.4 test')
    os.makedirs(os.path.join(fdir, '参考'))
    with open(os.path.join(fdir, '参考', '模板.docx'), 'wb') as f:
        f.write(b'x')
    r2 = A.api_contracts()
    ck('夹存在 → exists True + files≥2', r2['folder']['exists'] is True
       and len(r2['folder']['files']) >= 2)
    ck('files 带 kind/name/size_kb/mtime/sub',
       all(set(('kind', 'name', 'size_kb', 'mtime', 'sub')) <= set(x)
           for x in r2['folder']['files']))
    ck('子目录内文件 sub=参考',
       any(x['sub'] == '参考' and x['name'] == '模板.docx' for x in r2['folder']['files']))
    ck('调用后台账 mtime 未变（只读）', os.path.getmtime(ledger) == mtime0)

    print('=== ③ sheet 缺失：优雅降级 ===')
    ledger2 = os.path.join(TMP, '台账无合同sheet.xlsx')
    _make_ledger(ledger2, with_contract=False)
    os.environ['LEDGER_PATH'] = ledger2
    ck('read_rows guard：缺 sheet → []', L.read_rows(L.SHEET_CONTRACT) == [])
    r3 = A.api_contracts()
    ck('缺 sheet → ok + count 0 + items []',
       r3.get('ok') is True and r3['count'] == 0 and r3['items'] == []
       and r3['fee_total'] == 0)
    ck('缺 sheet 仍返回素材夹扫描', 'folder' in r3 and r3['folder']['exists'] is True)

    print('=== ④ 台账号字段直读（COLS 注册正确）===')
    os.environ['LEDGER_PATH'] = ledger
    rows = L.read_rows(L.SHEET_CONTRACT)
    ck('COLS_BY_SHEET 注册合同映射（且律师函映射未被覆盖）',
       L.COLS_BY_SHEET[L.SHEET_CONTRACT] is L.CONTRACT_COLS
       and L.COLS_BY_SHEET[L.SHEET_LETTER] is L.LETTER_COLS)
    ck('合同行第 10 列=备注 映射正确', rows[0].get('备注') in ('', None)
       and rows[2].get('交付日期') == '2026-10-05')

    os.environ.pop('LEDGER_PATH', None)
    shutil.rmtree(TMP, ignore_errors=True)
    print('\n=== 结果：%d/%d 通过 ===' % (len(total) - len(fails), len(total)))
    if fails:
        print('FAIL 项：')
        for f in fails:
            print('  - ' + f)
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
