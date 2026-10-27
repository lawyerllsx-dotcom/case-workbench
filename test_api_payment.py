# -*- coding: utf-8 -*-
r"""test_api_payment.py — 执行款回款记录接口回归（批 62，2026-10-13）

覆盖：
  · 清单 / 汇总（自建库只读）
  · 记录（confirm 门 dry-run / apply / 参数校验：日期·金额·夹名）
  · 删除（dry-run / apply；不回删台账行）
  · 同步到台账「执行款」sheet（**合成台账真跑全链**）：
    - 无 sheet → 自动受控建表（三行头）→ 追加 → db 回填 ledger_row/synced_at
    - 已同步 → action=none；再记一笔 → 直接追加（start_row 顺延）
    - 台账口径：案件/案号取匹配结果（打桩）
  · 安全闸：sync 的 case 越界拒绝

⚠️ 全隔离：`WORKBENCH_DB`/`WORKBENCH_BACKUP_DIR`/`LEDGER_PATH`/`LEDGER_BACKUP_ROOT`
   全部指临时目录（import app 前注入）；案件根打桩；**不碰真台账、真自建库**。
"""
import os
import shutil
import sys
import tempfile

TMP = tempfile.mkdtemp(prefix='wb_pay_test_')
os.environ['WORKBENCH_DB'] = os.path.join(TMP, 'wb.db')
os.environ['WORKBENCH_BACKUP_DIR'] = os.path.join(TMP, 'wbbak')
os.environ['LEDGER_PATH'] = os.path.join(TMP, 'fake_ledger.xlsx')
os.environ['LEDGER_BACKUP_ROOT'] = os.path.join(TMP, 'ledbak')

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

fails, total = [], []


def ck(name, cond):
    total.append(name)
    print(('  PASS ' if cond else '  FAIL ') + name)
    if not cond:
        fails.append(name)


def main():
    import openpyxl
    import app as A

    # 合成台账：只有「海大案件」（无「执行款」sheet——测自动建表路径）
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '海大案件'
    for i, h in enumerate(['序号', '案号', '案件'], start=1):
        ws.cell(row=3, column=i, value=h)
    wb.save(os.environ['LEDGER_PATH'])
    wb.close()

    tmp_root = os.path.join(TMP, 'cases')
    case = '执-测试案'
    os.makedirs(os.path.join(tmp_root, case))

    orig_root = A.case_ops.case_root
    orig_rows = A.case_ops.ledger_case_rows
    orig_match = A.case_ops.match_case_ledger
    try:
        A.case_ops.case_root = lambda: tmp_root
        A.case_ops.ledger_case_rows = lambda: []
        A.case_ops.match_case_ledger = lambda name, rows=None: {'matched': {
            '案件': '测试案（台账名）', '案号': '（2026）闽0000民初1号',
            'sheet': '海大案件', 'row': 4, 'match_score': 5}}

        print('=== ① 清单与记录 ===')
        r = A.api_payments(case)
        ck('空清单', r.get('ok') and r['items'] == [] and r['summary']['count'] == 0)
        r = A.api_payment_add({'folder': case, 'pay_date': '2026-10-01', 'amount': '50000'})
        ck('add 缺 confirm → dry-run 不落库', r.get('dry_run') is True and r['will_add']['amount'] == '50000')
        ck('dry-run 后清单仍空', A.api_payments(case)['summary']['count'] == 0)
        r = A.api_payment_add({'folder': case, 'pay_date': '2026-10-01', 'amount': '50000',
                               'source': '法院划付', 'note': '首笔', 'confirm': True})
        ck('add apply → id', r.get('ok') and r.get('id'))
        r = A.api_payment_add({'folder': case, 'pay_date': '2026/10/02', 'amount': '1', 'confirm': True})
        ck('坏日期 → 拒', 'error' in r)
        r = A.api_payment_add({'folder': case, 'pay_date': '2026-10-02', 'amount': 'abc', 'confirm': True})
        ck('坏金额 → 拒', 'error' in r)
        r = A.api_payment_add({'folder': '', 'pay_date': '2026-10-02', 'amount': '1', 'confirm': True})
        ck('空夹名 → 拒', 'error' in r)
        r = A.api_payment_add({'folder': case, 'pay_date': '2026-10-05', 'amount': '1000.5',
                               'source': '对方支付', 'confirm': True})
        ck('第二笔 apply → 清单 2 笔合计', r.get('ok')
           and A.api_payments(case)['summary'] == {'count': 2, 'total': 51000.5, 'pending': 2})

        print('=== ② 同步到台账（合成台账真跑：自动建表 → 追加 → 回填） ===')
        r = A.api_payment_sync({'folder': case, 'confirm': False})
        ck('sync dry-run：will_create + 起始 4 行 + 2 笔',
           r.get('ok') and r.get('dry_run') and r['plan']['will_create']
           and r['plan']['start_row'] == 4 and r['plan']['count'] == 2)
        ck('dry-run 后台账无「执行款」sheet', '执行款' not in openpyxl.load_workbook(
            os.environ['LEDGER_PATH']).sheetnames)
        r = A.api_payment_sync({'folder': case, 'confirm': True})
        ck('sync apply：建表 + 追加', r.get('ok') and r.get('action') == 'appended'
           and r.get('start_row') == 4 and r.get('count') == 2)
        wb2 = openpyxl.load_workbook(os.environ['LEDGER_PATH'])
        ws2 = wb2['执行款']
        ck('新表三行头：r1 标题 + r3 表头',
           ws2.cell(row=1, column=1).value == '执行款回款登记表'
           and [ws2.cell(row=3, column=c).value for c in range(1, 8)] == A.ledger_ops.PAYMENT_HEADER)
        row4 = [ws2.cell(row=4, column=c).value for c in range(1, 8)]
        row5 = [ws2.cell(row=5, column=c).value for c in range(1, 8)]
        wb2.close()
        ck('r4 数据对位（序号1/台账口径案件名/案号/日期/金额/来源）',
           row4[0] == 1 and row4[1] == '测试案（台账名）' and row4[2] == '（2026）闽0000民初1号'
           and row4[3] == '2026-10-01' and row4[4] == 50000 and row4[5] == '法院划付')
        ck('r5 序号续编 2', row5[0] == 2)
        rows_db = A.payment_ops.list_payments(case)
        ck('db 回填 ledger_row 4/5 + synced_at',
           [x['ledger_row'] for x in rows_db] == [4, 5] and all(x['synced_at'] for x in rows_db))
        ck('回填后 pending=0', A.api_payments(case)['summary']['pending'] == 0)
        r = A.api_payment_sync({'folder': case, 'confirm': True})
        ck('全同步后复跑 → action=none', r.get('action') == 'none')

        print('=== ③ 再记一笔 → 直接追加（sheet 已存在） ===')
        A.api_payment_add({'folder': case, 'pay_date': '2026-10-10', 'amount': '500',
                           'source': '其他', 'confirm': True})
        r = A.api_payment_sync({'folder': case, 'confirm': False})
        ck('dry-run：不建表 + 起始 6 行',
           (not r['plan']['will_create']) and r['plan']['start_row'] == 6 and r['plan']['count'] == 1)
        r = A.api_payment_sync({'folder': case, 'confirm': True})
        ck('apply 追加至 r6', r.get('ok') and r.get('start_row') == 6)

        print('=== ④ 删除（不回删台账行） ===')
        pid = A.payment_ops.list_payments(case)[0]['id']
        r = A.api_payment_delete({'id': pid})
        ck('delete 缺 confirm → dry-run', r.get('dry_run') is True)
        r = A.api_payment_delete({'id': pid, 'confirm': True})
        ck('delete apply → removed=1', r.get('ok') and r.get('removed') == 1)
        ck('db 剩 2 笔', len(A.payment_ops.list_payments(case)) == 2)
        wb3 = openpyxl.load_workbook(os.environ['LEDGER_PATH'])
        ck('台账行未被回删（仍 3 行数据）',
           wb3['执行款'].cell(row=7, column=1).value in (None, '')
           and wb3['执行款'].cell(row=6, column=1).value == 3)
        wb3.close()

        print('=== ⑤ 安全闸 ===')
        r = A.api_payment_sync({'folder': '..', 'confirm': True})
        ck('sync case 越界 → 拒', 'error' in r)
        r = A.api_payment_sync({'folder': '不存在的夹', 'confirm': False})
        ck('sync case 不存在 → 拒', 'error' in r)
    finally:
        A.case_ops.case_root = orig_root
        A.case_ops.ledger_case_rows = orig_rows
        A.case_ops.match_case_ledger = orig_match
        shutil.rmtree(TMP, ignore_errors=True)

    print('\n=== 结果：%d/%d 通过 ===' % (len(total) - len(fails), len(total)))
    if fails:
        print('FAIL 项：')
        for f in fails:
            print('  - ' + f)
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
